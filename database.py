"""
database.py - TQF System Database Layer
SQLite schema + CRUD helpers
"""

import json
import os
import sqlite3
from contextlib import contextmanager
from datetime import datetime


DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tqf_database.db")


@contextmanager
def get_conn():
    conn = sqlite3.connect(DB_PATH, timeout=10)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


SCHEMA = """
CREATE TABLE IF NOT EXISTS curricula (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    version         TEXT    NOT NULL UNIQUE,
    name_th         TEXT    DEFAULT '',
    effective_year  INTEGER DEFAULT 0,
    graduation_req  TEXT    DEFAULT '',
    created_at      TEXT    DEFAULT (datetime('now','localtime'))
);

CREATE TABLE IF NOT EXISTS courses (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    curriculum_id       INTEGER REFERENCES curricula(id),
    code                TEXT    NOT NULL,
    name_th             TEXT    NOT NULL,
    name_en             TEXT    NOT NULL DEFAULT '',
    credits_text        TEXT    DEFAULT '',
    credit_lecture      INTEGER DEFAULT 0,
    credit_lab          INTEGER DEFAULT 0,
    credit_self         INTEGER DEFAULT 0,
    prerequisite        TEXT    DEFAULT 'ไม่มี',
    course_type         TEXT    DEFAULT '',
    description_th      TEXT    DEFAULT '',
    description_en      TEXT    DEFAULT '',
    faculty             TEXT    DEFAULT '',
    department          TEXT    DEFAULT '',
    created_at          TEXT    DEFAULT (datetime('now','localtime')),
    UNIQUE(code, curriculum_id)
);

CREATE TABLE IF NOT EXISTS course_offerings (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    course_id           INTEGER NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
    curriculum_id       INTEGER NOT NULL REFERENCES curricula(id) ON DELETE CASCADE,
    semester            INTEGER NOT NULL,
    year                INTEGER NOT NULL,
    section_code        TEXT    NOT NULL,
    is_special          INTEGER NOT NULL DEFAULT 0,
    status              TEXT    NOT NULL DEFAULT 'active',
    source_type         TEXT    NOT NULL DEFAULT 'catalog',
    created_at          TEXT    DEFAULT (datetime('now','localtime')),
    updated_at          TEXT    DEFAULT (datetime('now','localtime')),
    UNIQUE(course_id, semester, year, section_code)
);

CREATE TABLE IF NOT EXISTS tqf3 (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    course_id           INTEGER NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
    offering_id         INTEGER REFERENCES course_offerings(id) ON DELETE SET NULL,
    semester            INTEGER NOT NULL,
    year                INTEGER NOT NULL,
    instructor_main     TEXT    DEFAULT '',
    instructors_json    TEXT    DEFAULT '[]',
    location            TEXT    DEFAULT '',
    objectives          TEXT    DEFAULT '',
    source_file         TEXT    DEFAULT '',
    source_type         TEXT    DEFAULT 'imported',
    is_special          INTEGER NOT NULL DEFAULT 0,
    imported_at         TEXT    DEFAULT (datetime('now','localtime')),
    UNIQUE(course_id, semester, year, is_special)
);

CREATE TABLE IF NOT EXISTS clos (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    tqf3_id             INTEGER NOT NULL REFERENCES tqf3(id) ON DELETE CASCADE,
    clo_number          INTEGER NOT NULL,
    description         TEXT    NOT NULL,
    teaching_strategy   TEXT    DEFAULT '',
    assessment_method   TEXT    DEFAULT '',
    indicator           TEXT    DEFAULT '',
    target_pct          REAL    DEFAULT 50.0,
    UNIQUE(tqf3_id, clo_number)
);

CREATE TABLE IF NOT EXISTS teaching_plan (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    tqf3_id             INTEGER NOT NULL REFERENCES tqf3(id) ON DELETE CASCADE,
    week                INTEGER NOT NULL,
    week_label          TEXT    DEFAULT '',
    llo_text            TEXT    DEFAULT '',
    topic               TEXT    NOT NULL,
    activities          TEXT    DEFAULT '',
    media               TEXT    DEFAULT '',
    assessment_tools    TEXT    DEFAULT '',
    hours_planned       REAL    NOT NULL DEFAULT 0,
    hours_theory        REAL    DEFAULT 0,
    hours_practice      REAL    DEFAULT 0,
    hours_self          REAL    DEFAULT 0,
    teaching_method     TEXT    DEFAULT ''
);

CREATE TABLE IF NOT EXISTS assessments (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    tqf3_id             INTEGER NOT NULL REFERENCES tqf3(id) ON DELETE CASCADE,
    name                TEXT    NOT NULL,
    full_score          REAL    DEFAULT 100,
    weight_pct          REAL    NOT NULL DEFAULT 0,
    clo_mapping         TEXT    DEFAULT '[]',
    assessment_period   TEXT    DEFAULT '',
    eval_criteria       TEXT    DEFAULT '',
    pass_threshold      REAL    DEFAULT 50.0
);

CREATE TABLE IF NOT EXISTS tqf5 (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    tqf3_id             INTEGER NOT NULL UNIQUE REFERENCES tqf3(id) ON DELETE CASCADE,
    registered_count    INTEGER DEFAULT 0,
    remaining_count     INTEGER DEFAULT 0,
    withdrawn_count     INTEGER DEFAULT 0,
    grade_dist_json     TEXT    DEFAULT '{}',
    abnormal_factors    TEXT    DEFAULT '',
    deviation_time      TEXT    DEFAULT '',
    deviation_method    TEXT    DEFAULT '',
    clo_results_json    TEXT    DEFAULT '[]',
    verification_method TEXT    DEFAULT '',
    resource_issues     TEXT    DEFAULT '',
    admin_issues        TEXT    DEFAULT '',
    student_eval_notes  TEXT    DEFAULT '',
    teacher_response    TEXT    DEFAULT '',
    improvement_prev    TEXT    DEFAULT '',
    improvement_next    TEXT    DEFAULT '',
    improvement_resp    TEXT    DEFAULT '',
    suggestions         TEXT    DEFAULT '',
    last_updated        TEXT    DEFAULT (datetime('now','localtime'))
);

CREATE TABLE IF NOT EXISTS student_grades (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    tqf5_id             INTEGER NOT NULL REFERENCES tqf5(id) ON DELETE CASCADE,
    student_id          TEXT    NOT NULL,
    student_name        TEXT    NOT NULL,
    section             TEXT    DEFAULT '',
    score_components    TEXT    DEFAULT '{}',
    total_score         REAL    DEFAULT 0,
    grade               TEXT    NOT NULL,
    UNIQUE(tqf5_id, student_id)
);

CREATE TABLE IF NOT EXISTS teaching_actual (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    tqf5_id             INTEGER NOT NULL REFERENCES tqf5(id) ON DELETE CASCADE,
    topic               TEXT    NOT NULL,
    hours_planned       REAL    DEFAULT 0,
    hours_actual        REAL    DEFAULT 0,
    deviation_reason    TEXT    DEFAULT ''
);

CREATE TABLE IF NOT EXISTS plos (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    curriculum_id   INTEGER NOT NULL REFERENCES curricula(id) ON DELETE CASCADE,
    plo_number      INTEGER NOT NULL,
    plo_code        TEXT    DEFAULT '',
    category        TEXT    DEFAULT '',
    description     TEXT    NOT NULL DEFAULT '',
    UNIQUE(curriculum_id, plo_number)
);

CREATE TABLE IF NOT EXISTS ylos (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    curriculum_id       INTEGER NOT NULL REFERENCES curricula(id) ON DELETE CASCADE,
    year_number         INTEGER NOT NULL,
    title               TEXT    NOT NULL DEFAULT '',
    indicators          TEXT    DEFAULT '',
    assessment_methods  TEXT    DEFAULT '',
    plo_mapping         TEXT    DEFAULT '[]',
    UNIQUE(curriculum_id, year_number)
);

CREATE TABLE IF NOT EXISTS course_clos (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    course_id           INTEGER NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
    clo_number          INTEGER NOT NULL,
    description         TEXT    NOT NULL DEFAULT '',
    domain              TEXT    DEFAULT '',
    teaching_strategy   TEXT    DEFAULT '',
    assessment_method   TEXT    DEFAULT '',
    pass_threshold_pct  REAL    DEFAULT 50.0,
    plo_mapping         TEXT    DEFAULT '[]',
    UNIQUE(course_id, clo_number)
);

CREATE TABLE IF NOT EXISTS course_teaching_plan (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    course_id           INTEGER NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
    seq                 INTEGER DEFAULT 0,
    week                INTEGER NOT NULL DEFAULT 0,
    week_label          TEXT    DEFAULT '',
    llo_text            TEXT    DEFAULT '',
    topic               TEXT    NOT NULL DEFAULT '',
    activities          TEXT    DEFAULT '',
    teaching_method     TEXT    DEFAULT '',
    media               TEXT    DEFAULT '',
    assessment_tools    TEXT    DEFAULT '',
    hours_planned       REAL    DEFAULT 0,
    hours_theory        REAL    DEFAULT 0,
    hours_practice      REAL    DEFAULT 0,
    hours_self          REAL    DEFAULT 0
);

CREATE TABLE IF NOT EXISTS course_assessments (
    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
    course_id           INTEGER NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
    seq                 INTEGER DEFAULT 0,
    name                TEXT    NOT NULL,
    full_score          REAL    DEFAULT 100,
    weight_pct          REAL    NOT NULL DEFAULT 0,
    clo_mapping         TEXT    DEFAULT '[]',
    assessment_period   TEXT    DEFAULT '',
    eval_criteria       TEXT    DEFAULT '',
    pass_threshold      REAL    DEFAULT 50.0
);

CREATE TABLE IF NOT EXISTS course_resources (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    course_id       INTEGER NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
    seq             INTEGER DEFAULT 0,
    resource_type   TEXT    DEFAULT '',
    citation_text   TEXT    NOT NULL DEFAULT '',
    url             TEXT    DEFAULT '',
    note            TEXT    DEFAULT ''
);

CREATE TABLE IF NOT EXISTS tqf3_staff (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    tqf3_id         INTEGER NOT NULL REFERENCES tqf3(id) ON DELETE CASCADE,
    role            TEXT    NOT NULL DEFAULT '',
    seq             INTEGER DEFAULT 0,
    name            TEXT    NOT NULL DEFAULT ''
);
"""


def init_db():
    with get_conn() as conn:
        conn.executescript(SCHEMA)

        def ensure_column(table_name: str, column_name: str, definition: str):
            cols = [r[1] for r in conn.execute(f"PRAGMA table_info({table_name})").fetchall()]
            if column_name not in cols:
                conn.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_name} {definition}")
                print(f"[DB] Migration: added {table_name}.{column_name}")

        ensure_column("tqf3", "is_special", "INTEGER NOT NULL DEFAULT 0")
        ensure_column("tqf3", "source_type", "TEXT DEFAULT 'imported'")
        ensure_column("tqf3", "offering_id", "INTEGER REFERENCES course_offerings(id)")
        ensure_column("courses", "curriculum_id", "INTEGER REFERENCES curricula(id)")
        ensure_column("teaching_plan", "week_label", "TEXT DEFAULT ''")
        ensure_column("teaching_plan", "llo_text", "TEXT DEFAULT ''")
        ensure_column("teaching_plan", "activities", "TEXT DEFAULT ''")
        ensure_column("teaching_plan", "media", "TEXT DEFAULT ''")
        ensure_column("teaching_plan", "assessment_tools", "TEXT DEFAULT ''")
        ensure_column("teaching_plan", "hours_theory", "REAL DEFAULT 0")
        ensure_column("teaching_plan", "hours_practice", "REAL DEFAULT 0")
        ensure_column("teaching_plan", "hours_self", "REAL DEFAULT 0")
        ensure_column("assessments", "full_score", "REAL DEFAULT 100")
        ensure_column("assessments", "assessment_period", "TEXT DEFAULT ''")
        ensure_column("assessments", "eval_criteria", "TEXT DEFAULT ''")
        ensure_column("assessments", "pass_threshold", "REAL DEFAULT 50.0")
        ensure_column("course_assessments", "full_score", "REAL DEFAULT 100")
        ensure_column("course_assessments", "assessment_period", "TEXT DEFAULT ''")
        ensure_column("course_assessments", "eval_criteria", "TEXT DEFAULT ''")
        ensure_column("course_assessments", "pass_threshold", "REAL DEFAULT 50.0")
        ensure_column("plos", "plo_code", "TEXT DEFAULT ''")
        ensure_column("curricula", "graduation_req", "TEXT DEFAULT ''")

        def _index_columns(table_name: str):
            result = []
            for row in conn.execute(f"PRAGMA index_list({table_name})").fetchall():
                if not row[2]:
                    continue
                idx_name = row[1]
                cols = [
                    info[2]
                    for info in conn.execute(f"PRAGMA index_info('{idx_name}')").fetchall()
                ]
                result.append(cols)
            return result

        def ensure_tqf3_unique_shape():
            expected = ["course_id", "semester", "year", "is_special"]
            unique_indexes = _index_columns("tqf3")
            if expected in unique_indexes:
                return

            conn.execute("PRAGMA foreign_keys=OFF")
            conn.execute(
                """
                CREATE TABLE tqf3__new (
                    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
                    course_id           INTEGER NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
                    offering_id         INTEGER REFERENCES course_offerings(id) ON DELETE SET NULL,
                    semester            INTEGER NOT NULL,
                    year                INTEGER NOT NULL,
                    instructor_main     TEXT    DEFAULT '',
                    instructors_json    TEXT    DEFAULT '[]',
                    location            TEXT    DEFAULT '',
                    objectives          TEXT    DEFAULT '',
                    source_file         TEXT    DEFAULT '',
                    source_type         TEXT    DEFAULT 'imported',
                    is_special          INTEGER NOT NULL DEFAULT 0,
                    imported_at         TEXT    DEFAULT (datetime('now','localtime')),
                    UNIQUE(course_id, semester, year, is_special)
                )
                """
            )
            conn.execute(
                """
                INSERT INTO tqf3__new (
                    id, course_id, offering_id, semester, year, instructor_main,
                    instructors_json, location, objectives, source_file,
                    source_type, is_special, imported_at
                )
                SELECT
                    id, course_id, offering_id, semester, year,
                    COALESCE(instructor_main, ''),
                    COALESCE(instructors_json, '[]'),
                    COALESCE(location, ''),
                    COALESCE(objectives, ''),
                    COALESCE(source_file, ''),
                    COALESCE(source_type, 'imported'),
                    COALESCE(is_special, 0),
                    COALESCE(imported_at, datetime('now','localtime'))
                FROM tqf3
                """
            )
            conn.execute("DROP TABLE tqf3")
            conn.execute("ALTER TABLE tqf3__new RENAME TO tqf3")
            conn.execute("PRAGMA foreign_keys=ON")
            print("[DB] Migration: rebuilt tqf3 unique constraint to include is_special")

        ensure_tqf3_unique_shape()
        ensure_column("tqf3", "offering_id", "INTEGER REFERENCES course_offerings(id)")

        def ensure_course_offerings_backfill():
            rows = conn.execute(
                """
                SELECT
                    t.id AS tqf3_id,
                    t.course_id,
                    t.offering_id,
                    t.semester,
                    t.year,
                    COALESCE(t.is_special, 0) AS is_special,
                    COALESCE(t.source_type, 'imported') AS source_type,
                    COALESCE(c.curriculum_id, 0) AS curriculum_id
                FROM tqf3 t
                JOIN courses c ON c.id = t.course_id
                """
            ).fetchall()
            if not rows:
                return

            touched = 0
            for row in rows:
                section_code = "P01" if row["is_special"] else "N01"
                conn.execute(
                    """
                    INSERT INTO course_offerings (
                        course_id, curriculum_id, semester, year,
                        section_code, is_special, status, source_type
                    )
                    VALUES (?,?,?,?,?,?,?,?)
                    ON CONFLICT(course_id, semester, year, section_code) DO UPDATE SET
                        curriculum_id=excluded.curriculum_id,
                        is_special=excluded.is_special,
                        updated_at=datetime('now','localtime')
                    """,
                    (
                        row["course_id"],
                        row["curriculum_id"],
                        row["semester"],
                        row["year"],
                        section_code,
                        row["is_special"],
                        "active",
                        row["source_type"] or "catalog",
                    ),
                )
                offering = conn.execute(
                    """
                    SELECT id FROM course_offerings
                    WHERE course_id=? AND semester=? AND year=? AND section_code=?
                    """,
                    (row["course_id"], row["semester"], row["year"], section_code),
                ).fetchone()
                if offering and row["offering_id"] != offering["id"]:
                    conn.execute(
                        "UPDATE tqf3 SET offering_id=? WHERE id=?",
                        (offering["id"], row["tqf3_id"]),
                    )
                    touched += 1

            if touched:
                print(f"[DB] Migration: linked {touched} tqf3 row(s) to course_offerings")

        ensure_course_offerings_backfill()

    print(f"[DB] Initialized: {DB_PATH}")


def upsert_course(code, name_th, name_en="", credits_text="",
                  credit_lecture=0, credit_lab=0, credit_self=0,
                  prerequisite="ไม่มี", course_type="",
                  description_th="", description_en="",
                  faculty="", department="",
                  curriculum_id=None) -> int:
    with get_conn() as conn:
        if curriculum_id is None:
            existing = conn.execute(
                "SELECT id, curriculum_id FROM courses WHERE code=? ORDER BY curriculum_id DESC LIMIT 1",
                (code,),
            ).fetchone()
            if existing:
                conn.execute(
                    """
                    UPDATE courses SET
                        name_th=CASE WHEN ?!='' THEN ? ELSE name_th END,
                        name_en=CASE WHEN ?!='' THEN ? ELSE name_en END,
                        credits_text=CASE WHEN ?!='' THEN ? ELSE credits_text END,
                        credit_lecture=CASE WHEN ?>0 THEN ? ELSE credit_lecture END,
                        credit_lab=CASE WHEN ?>0 THEN ? ELSE credit_lab END,
                        credit_self=CASE WHEN ?>0 THEN ? ELSE credit_self END,
                        faculty=CASE WHEN ?!='' THEN ? ELSE faculty END,
                        department=CASE WHEN ?!='' THEN ? ELSE department END
                    WHERE id=?
                    """,
                    (
                        name_th, name_th, name_en, name_en, credits_text, credits_text,
                        credit_lecture, credit_lecture, credit_lab, credit_lab,
                        credit_self, credit_self, faculty, faculty, department, department,
                        existing["id"],
                    ),
                )
                return existing["id"]

        conn.execute(
            """
            INSERT INTO courses (
                code, name_th, name_en, credits_text,
                credit_lecture, credit_lab, credit_self, prerequisite,
                course_type, description_th, description_en,
                faculty, department, curriculum_id
            )
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(code, curriculum_id) DO UPDATE SET
                name_th=excluded.name_th,
                name_en=excluded.name_en,
                credits_text=excluded.credits_text,
                credit_lecture=excluded.credit_lecture,
                credit_lab=excluded.credit_lab,
                credit_self=excluded.credit_self,
                prerequisite=excluded.prerequisite,
                course_type=excluded.course_type,
                description_th=excluded.description_th,
                description_en=excluded.description_en,
                faculty=excluded.faculty,
                department=excluded.department
            """,
            (
                code, name_th, name_en, credits_text,
                credit_lecture, credit_lab, credit_self, prerequisite,
                course_type, description_th, description_en,
                faculty, department, curriculum_id,
            ),
        )

        row = conn.execute(
            "SELECT id FROM courses WHERE code=? AND curriculum_id IS ?",
            (code, curriculum_id),
        ).fetchone()
        if not row:
            row = conn.execute("SELECT id FROM courses WHERE code=?", (code,)).fetchone()
        return row["id"]


def get_all_courses():
    with get_conn() as conn:
        return [dict(r) for r in conn.execute("SELECT * FROM courses ORDER BY code").fetchall()]


def get_course_by_code(code, curriculum_id=None):
    with get_conn() as conn:
        if curriculum_id is not None:
            row = conn.execute(
                "SELECT * FROM courses WHERE code=? AND curriculum_id=?",
                (code, curriculum_id),
            ).fetchone()
        else:
            row = conn.execute(
                "SELECT * FROM courses WHERE code=? ORDER BY curriculum_id DESC LIMIT 1",
                (code,),
            ).fetchone()
        return dict(row) if row else None


def upsert_course_offering(course_id, semester, year, curriculum_id=None,
                           section_code="", is_special=None,
                           status="active", source_type="catalog") -> int:
    with get_conn() as conn:
        if curriculum_id is None:
            course = conn.execute(
                "SELECT curriculum_id FROM courses WHERE id=?",
                (course_id,),
            ).fetchone()
            curriculum_id = course["curriculum_id"] if course else None
        if curriculum_id is None:
            raise ValueError(f"Course {course_id} has no curriculum_id")

        if is_special is None:
            is_special = str(section_code or "").upper().startswith("P")
        is_special = int(bool(is_special))

        section_code = (section_code or "").strip().upper()
        if not section_code:
            section_code = "P01" if is_special else "N01"

        conn.execute(
            """
            INSERT INTO course_offerings (
                course_id, curriculum_id, semester, year,
                section_code, is_special, status, source_type
            )
            VALUES (?,?,?,?,?,?,?,?)
            ON CONFLICT(course_id, semester, year, section_code) DO UPDATE SET
                curriculum_id=excluded.curriculum_id,
                is_special=excluded.is_special,
                status=excluded.status,
                source_type=excluded.source_type,
                updated_at=datetime('now','localtime')
            """,
            (
                course_id,
                curriculum_id,
                semester,
                year,
                section_code,
                is_special,
                status,
                source_type,
            ),
        )
        row = conn.execute(
            """
            SELECT id FROM course_offerings
            WHERE course_id=? AND semester=? AND year=? AND section_code=?
            """,
            (course_id, semester, year, section_code),
        ).fetchone()
        return row["id"]


def get_course_offerings(course_id=None, curriculum_id=None, semester=None, year=None):
    where = []
    params = []
    if course_id is not None:
        where.append("o.course_id=?")
        params.append(course_id)
    if curriculum_id is not None:
        where.append("o.curriculum_id=?")
        params.append(curriculum_id)
    if semester is not None:
        where.append("o.semester=?")
        params.append(semester)
    if year is not None:
        where.append("o.year=?")
        params.append(year)

    where_sql = f"WHERE {' AND '.join(where)}" if where else ""
    with get_conn() as conn:
        rows = conn.execute(
            f"""
            SELECT
                o.*,
                c.code,
                c.name_th,
                c.course_type,
                cu.version AS curriculum_version,
                t.id AS tqf3_id,
                COALESCE(t.source_type, '') AS tqf3_source_type
            FROM course_offerings o
            JOIN courses c ON c.id = o.course_id
            LEFT JOIN curricula cu ON cu.id = o.curriculum_id
            LEFT JOIN tqf3 t ON t.offering_id = o.id
            {where_sql}
            ORDER BY o.year DESC, o.semester DESC, o.section_code, c.code
            """,
            params,
        ).fetchall()
        return [dict(r) for r in rows]


def upsert_tqf3(course_id, semester, year, instructor_main="",
                instructors=None, location="", objectives="",
                source_file="", is_special=0, offering_id=None,
                section_code="") -> int:
    instructors_json = json.dumps(instructors or [], ensure_ascii=False)
    is_special = int(bool(is_special))
    with get_conn() as conn:
        if offering_id is None:
            course = conn.execute(
                "SELECT curriculum_id FROM courses WHERE id=?",
                (course_id,),
            ).fetchone()
            if not course or course["curriculum_id"] is None:
                raise ValueError(f"Course {course_id} has no curriculum_id")
            section_code = (section_code or "").strip().upper() or ("P01" if is_special else "N01")
            conn.execute(
                """
                INSERT INTO course_offerings (
                    course_id, curriculum_id, semester, year,
                    section_code, is_special, status, source_type
                )
                VALUES (?,?,?,?,?,?,?,?)
                ON CONFLICT(course_id, semester, year, section_code) DO UPDATE SET
                    curriculum_id=excluded.curriculum_id,
                    is_special=excluded.is_special,
                    updated_at=datetime('now','localtime')
                """,
                (
                    course_id,
                    course["curriculum_id"],
                    semester,
                    year,
                    section_code,
                    is_special,
                    "active",
                    "imported" if source_file else "catalog",
                ),
            )
            offering = conn.execute(
                """
                SELECT id FROM course_offerings
                WHERE course_id=? AND semester=? AND year=? AND section_code=?
                """,
                (course_id, semester, year, section_code),
            ).fetchone()
            offering_id = offering["id"]

        conn.execute(
            """
            INSERT INTO tqf3 (
                course_id, offering_id, semester, year, instructor_main,
                instructors_json, location, objectives, source_file, is_special
            )
            VALUES (?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(course_id, semester, year, is_special) DO UPDATE SET
                offering_id=excluded.offering_id,
                instructor_main=excluded.instructor_main,
                instructors_json=excluded.instructors_json,
                location=excluded.location,
                objectives=excluded.objectives,
                source_file=excluded.source_file,
                imported_at=datetime('now','localtime')
            """,
            (
                course_id, offering_id, semester, year, instructor_main,
                instructors_json, location, objectives, source_file, is_special,
            ),
        )
        row = conn.execute(
            """
            SELECT id FROM tqf3
            WHERE course_id=? AND semester=? AND year=? AND is_special=?
            ORDER BY CASE WHEN offering_id=? THEN 0 ELSE 1 END, id
            LIMIT 1
            """,
            (course_id, semester, year, is_special, offering_id),
        ).fetchone()
        return row["id"]


def get_tqf3(tqf3_id):
    with get_conn() as conn:
        row = conn.execute(
            """
            SELECT t.*, c.code, c.name_th
            FROM tqf3 t
            JOIN courses c ON c.id=t.course_id
            LEFT JOIN course_offerings o ON o.id=t.offering_id
            WHERE t.id=?
            """,
            (tqf3_id,),
        ).fetchone()
        if not row:
            return None
        data = dict(row)
        data["instructors"] = json.loads(data.get("instructors_json", "[]"))
        data["staff"] = get_tqf3_staff(tqf3_id)
        return data


def get_all_tqf3(course_id=None):
    with get_conn() as conn:
        if course_id:
            rows = conn.execute(
                """
                SELECT t.*, c.code, c.name_th, o.section_code
                FROM tqf3 t
                JOIN courses c ON c.id=t.course_id
                LEFT JOIN course_offerings o ON o.id=t.offering_id
                WHERE t.course_id=?
                ORDER BY t.year DESC, t.semester DESC
                """,
                (course_id,),
            ).fetchall()
        else:
            rows = conn.execute(
                """
                SELECT t.*, c.code, c.name_th, o.section_code
                FROM tqf3 t
                JOIN courses c ON c.id=t.course_id
                LEFT JOIN course_offerings o ON o.id=t.offering_id
                ORDER BY t.year DESC, t.semester DESC
                """
            ).fetchall()
        return [dict(r) for r in rows]


def replace_clos(tqf3_id, clos_list):
    with get_conn() as conn:
        conn.execute("DELETE FROM clos WHERE tqf3_id=?", (tqf3_id,))
        for c in clos_list:
            conn.execute(
                """
                INSERT INTO clos (
                    tqf3_id, clo_number, description,
                    teaching_strategy, assessment_method, indicator, target_pct
                )
                VALUES (?,?,?,?,?,?,?)
                """,
                (
                    tqf3_id,
                    c.get("clo_number", 0),
                    c.get("description", ""),
                    c.get("teaching_strategy", ""),
                    c.get("assessment_method", ""),
                    c.get("indicator", ""),
                    c.get("target_pct", 50.0),
                ),
            )


def get_clos(tqf3_id):
    with get_conn() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM clos WHERE tqf3_id=? ORDER BY clo_number",
            (tqf3_id,),
        ).fetchall()]


def replace_teaching_plan(tqf3_id, plan_list):
    with get_conn() as conn:
        conn.execute("DELETE FROM teaching_plan WHERE tqf3_id=?", (tqf3_id,))
        for p in plan_list:
            conn.execute(
                """
                INSERT INTO teaching_plan (
                    tqf3_id, week, week_label, llo_text, topic,
                    activities, media, assessment_tools,
                    hours_planned, hours_theory, hours_practice, hours_self,
                    teaching_method
                )
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    tqf3_id,
                    p.get("week", 0),
                    p.get("week_label", ""),
                    p.get("llo_text", ""),
                    p.get("topic", ""),
                    p.get("activities", ""),
                    p.get("media", ""),
                    p.get("assessment_tools", ""),
                    p.get("hours_planned", 0),
                    p.get("hours_theory", 0),
                    p.get("hours_practice", 0),
                    p.get("hours_self", 0),
                    p.get("teaching_method", ""),
                ),
            )


def get_teaching_plan(tqf3_id):
    with get_conn() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM teaching_plan WHERE tqf3_id=? ORDER BY week, id",
            (tqf3_id,),
        ).fetchall()]


def replace_assessments(tqf3_id, assessments_list):
    with get_conn() as conn:
        conn.execute("DELETE FROM assessments WHERE tqf3_id=?", (tqf3_id,))
        for a in assessments_list:
            conn.execute(
                """
                INSERT INTO assessments (
                    tqf3_id, name, full_score, weight_pct, clo_mapping,
                    assessment_period, eval_criteria, pass_threshold
                )
                VALUES (?,?,?,?,?,?,?,?)
                """,
                (
                    tqf3_id,
                    a.get("name", ""),
                    a.get("full_score", 100),
                    a.get("weight_pct", 0),
                    json.dumps(a.get("clo_mapping", []), ensure_ascii=False),
                    a.get("assessment_period", ""),
                    a.get("eval_criteria", ""),
                    a.get("pass_threshold", 50.0),
                ),
            )


def get_assessments(tqf3_id):
    with get_conn() as conn:
        rows = [dict(r) for r in conn.execute(
            "SELECT * FROM assessments WHERE tqf3_id=? ORDER BY id",
            (tqf3_id,),
        ).fetchall()]
        for row in rows:
            row["clo_mapping"] = json.loads(row.get("clo_mapping", "[]"))
        return rows


def get_or_create_tqf5(tqf3_id) -> int:
    with get_conn() as conn:
        row = conn.execute("SELECT id FROM tqf5 WHERE tqf3_id=?", (tqf3_id,)).fetchone()
        if row:
            return row["id"]
        cur = conn.execute("INSERT INTO tqf5 (tqf3_id) VALUES (?)", (tqf3_id,))
        return cur.lastrowid


def get_tqf5(tqf3_id):
    with get_conn() as conn:
        row = conn.execute("SELECT * FROM tqf5 WHERE tqf3_id=?", (tqf3_id,)).fetchone()
        if not row:
            return None
        data = dict(row)
        data["grade_dist"] = json.loads(data.get("grade_dist_json", "{}"))
        data["clo_results"] = json.loads(data.get("clo_results_json", "[]"))
        return data


def update_tqf5(tqf5_id, **kwargs):
    if "grade_dist" in kwargs:
        kwargs["grade_dist_json"] = json.dumps(kwargs.pop("grade_dist"), ensure_ascii=False)
    if "clo_results" in kwargs:
        kwargs["clo_results_json"] = json.dumps(kwargs.pop("clo_results"), ensure_ascii=False)
    if not kwargs:
        return
    kwargs["last_updated"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cols = ", ".join(f"{key}=?" for key in kwargs)
    vals = list(kwargs.values()) + [tqf5_id]
    with get_conn() as conn:
        conn.execute(f"UPDATE tqf5 SET {cols} WHERE id=?", vals)


def replace_student_grades(tqf5_id, grades_list):
    with get_conn() as conn:
        conn.execute("DELETE FROM student_grades WHERE tqf5_id=?", (tqf5_id,))
        for g in grades_list:
            conn.execute(
                """
                INSERT INTO student_grades (
                    tqf5_id, student_id, student_name, section,
                    score_components, total_score, grade
                )
                VALUES (?,?,?,?,?,?,?)
                """,
                (
                    tqf5_id,
                    g.get("student_id", ""),
                    g.get("student_name", ""),
                    g.get("section", ""),
                    json.dumps(g.get("score_components", {}), ensure_ascii=False),
                    g.get("total_score", 0),
                    g.get("grade", ""),
                ),
            )


def get_student_grades(tqf5_id):
    with get_conn() as conn:
        rows = [dict(r) for r in conn.execute(
            "SELECT * FROM student_grades WHERE tqf5_id=? ORDER BY student_id",
            (tqf5_id,),
        ).fetchall()]
        for row in rows:
            row["score_components"] = json.loads(row.get("score_components", "{}"))
        return rows


def compute_grade_stats(tqf5_id) -> dict:
    grades = get_student_grades(tqf5_id)
    grade_order = [
        "A", "B+", "B", "C+", "C", "D+", "D", "E",
        "ไม่สมบูรณ์ (I)", "ผ่าน (P,S)", "ไม่ผ่าน (U)", "W",
    ]
    dist = {grade: 0 for grade in grade_order}
    for student in grades:
        grade = student["grade"]
        if grade in dist:
            dist[grade] += 1
        elif grade == "I":
            dist["ไม่สมบูรณ์ (I)"] += 1
        elif grade in ("P", "S"):
            dist["ผ่าน (P,S)"] += 1
        elif grade == "U":
            dist["ไม่ผ่าน (U)"] += 1
    total = len(grades)
    withdrawn = dist.get("W", 0)
    return {
        "registered": total,
        "withdrawn": withdrawn,
        "remaining": total - withdrawn,
        "dist": dist,
        "pct": {
            grade: round(count / total * 100, 2) if total > 0 else 0.0
            for grade, count in dist.items()
        },
        "total": total,
    }


def replace_teaching_actual(tqf5_id, actual_list):
    with get_conn() as conn:
        conn.execute("DELETE FROM teaching_actual WHERE tqf5_id=?", (tqf5_id,))
        for item in actual_list:
            conn.execute(
                """
                INSERT INTO teaching_actual (
                    tqf5_id, topic, hours_planned, hours_actual, deviation_reason
                )
                VALUES (?,?,?,?,?)
                """,
                (
                    tqf5_id,
                    item.get("topic", ""),
                    item.get("hours_planned", 0),
                    item.get("hours_actual", 0),
                    item.get("deviation_reason", ""),
                ),
            )


def get_teaching_actual(tqf5_id):
    with get_conn() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM teaching_actual WHERE tqf5_id=? ORDER BY id",
            (tqf5_id,),
        ).fetchall()]


def get_dashboard_summary():
    with get_conn() as conn:
        total_courses = conn.execute("SELECT COUNT(*) FROM courses").fetchone()[0]
        total_tqf3 = conn.execute("SELECT COUNT(*) FROM tqf3").fetchone()[0]
        total_tqf5 = conn.execute("SELECT COUNT(*) FROM tqf5").fetchone()[0]
        total_students = conn.execute("SELECT COUNT(*) FROM student_grades").fetchone()[0]
        recent = conn.execute(
            """
            SELECT c.code, c.name_th, t.semester, t.year,
                   t5.registered_count, t5.grade_dist_json
            FROM tqf3 t
            JOIN courses c ON c.id = t.course_id
            LEFT JOIN tqf5 t5 ON t5.tqf3_id = t.id
            ORDER BY t.year DESC, t.semester DESC
            LIMIT 10
            """
        ).fetchall()
        return {
            "total_courses": total_courses,
            "total_tqf3": total_tqf3,
            "total_tqf5": total_tqf5,
            "total_students": total_students,
            "recent": [dict(r) for r in recent],
        }


def get_course_history(course_code):
    with get_conn() as conn:
        rows = conn.execute(
            """
            SELECT t.semester, t.year, t5.registered_count, t5.remaining_count,
                   t5.withdrawn_count, t5.grade_dist_json, t5.last_updated
            FROM tqf3 t
            JOIN courses c ON c.id = t.course_id
            LEFT JOIN tqf5 t5 ON t5.tqf3_id = t.id
            WHERE c.code = ?
            ORDER BY t.year DESC, t.semester DESC
            """,
            (course_code,),
        ).fetchall()
        result = []
        for row in rows:
            data = dict(row)
            data["grade_dist"] = json.loads(data.get("grade_dist_json") or "{}")
            result.append(data)
        return result


def upsert_curriculum(version: str, name_th: str = "", effective_year: int = 0) -> int:
    with get_conn() as conn:
        conn.execute(
            """
            INSERT INTO curricula (version, name_th, effective_year)
            VALUES (?, ?, ?)
            ON CONFLICT(version) DO UPDATE SET
                name_th=excluded.name_th,
                effective_year=excluded.effective_year
            """,
            (version, name_th, effective_year),
        )
        row = conn.execute("SELECT id FROM curricula WHERE version=?", (version,)).fetchone()
        return row["id"]


def get_curriculum_by_version(version: str):
    with get_conn() as conn:
        return conn.execute("SELECT * FROM curricula WHERE version=?", (version,)).fetchone()


def export_to_excel(output_path: str = None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if output_path is None:
        output_path = os.path.join(os.path.dirname(DB_PATH), "tqf_database_view.xlsx")

    header_fill = PatternFill("solid", start_color="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    row_a = PatternFill("solid", start_color="DEEAF1")
    row_b = PatternFill("solid", start_color="FFFFFF")
    normal = Font(name="Arial", size=10)

    def write_sheet(ws, headers, rows, widths):
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(1, col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22
        for row_idx, row in enumerate(rows, 2):
            ws.append(list(row))
            fill = row_a if row_idx % 2 == 0 else row_b
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row_idx, col_idx)
                cell.fill = fill
                cell.font = normal
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.row_dimensions[row_idx].height = 20
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"

    wb = Workbook()
    with get_conn() as conn:
        ws1 = wb.active
        ws1.title = "ภาพรวมรายวิชา"
        rows1 = conn.execute(
            """
            SELECT c.code, c.name_th, c.credits_text,
                   t.semester||'/'||t.year,
                   t.instructor_main,
                   (SELECT COUNT(*) FROM clos cl WHERE cl.tqf3_id=t.id),
                   COALESCE(t5.registered_count, '-'),
                   COALESCE(t5.remaining_count, '-'),
                   COALESCE(t5.withdrawn_count, '-')
            FROM courses c
            JOIN tqf3 t ON t.course_id = c.id
            LEFT JOIN tqf5 t5 ON t5.tqf3_id = t.id
            ORDER BY c.code
            """
        ).fetchall()
        write_sheet(
            ws1,
            ["รหัสวิชา", "ชื่อวิชา", "หน่วยกิต", "ภาค/ปี", "อาจารย์ผู้รับผิดชอบ", "CLOs", "ลงทะเบียน", "คงอยู่", "ถอน"],
            rows1,
            [12, 30, 10, 8, 28, 7, 10, 8, 8],
        )

        ws2 = wb.create_sheet("CLOs")
        rows2 = conn.execute(
            """
            SELECT c.code, c.name_th, cl.clo_number, cl.description
            FROM clos cl
            JOIN tqf3 t ON t.id = cl.tqf3_id
            JOIN courses c ON c.id = t.course_id
            ORDER BY c.code, cl.clo_number
            """
        ).fetchall()
        write_sheet(ws2, ["รหัสวิชา", "ชื่อวิชา", "CLO#", "คำอธิบาย"], rows2, [12, 24, 6, 70])

        ws3 = wb.create_sheet("การกระจายเกรด")
        grade_keys = ["A", "B+", "B", "C+", "C", "D+", "D", "E", "W"]
        rows3 = conn.execute(
            """
            SELECT c.code, c.name_th,
                   t5.registered_count, t5.remaining_count, t5.withdrawn_count,
                   t5.grade_dist_json
            FROM tqf5 t5
            JOIN tqf3 t ON t.id = t5.tqf3_id
            JOIN courses c ON c.id = t.course_id
            ORDER BY c.code
            """
        ).fetchall()
        data3 = []
        for row in rows3:
            dist = json.loads(row[5]) if row[5] else {}
            data3.append(list(row[:5]) + [dist.get(key, 0) for key in grade_keys])
        write_sheet(
            ws3,
            ["รหัสวิชา", "ชื่อวิชา", "ลงทะเบียน", "คงอยู่", "ถอน"] + grade_keys,
            data3,
            [12, 28, 10, 8, 8] + [6] * len(grade_keys),
        )

        ws4 = wb.create_sheet("รายชื่อนักศึกษา")
        rows4 = conn.execute(
            """
            SELECT c.code, sg.student_id, sg.student_name, sg.total_score, sg.grade
            FROM student_grades sg
            JOIN tqf5 t5 ON t5.id = sg.tqf5_id
            JOIN tqf3 t ON t.id = t5.tqf3_id
            JOIN courses c ON c.id = t.course_id
            ORDER BY c.code, sg.student_id
            """
        ).fetchall()
        write_sheet(
            ws4,
            ["รหัสวิชา", "รหัสนักศึกษา", "ชื่อ-สกุล", "คะแนนรวม", "เกรด"],
            rows4,
            [12, 18, 36, 12, 8],
        )

    wb.save(output_path)
    print(f"[Excel] Export -> {output_path}")
    return output_path


def upsert_plo(curriculum_id: int, plo_number: int,
               category: str = "", description: str = "",
               plo_code: str = "") -> int:
    with get_conn() as conn:
        conn.execute(
            """
            INSERT INTO plos (curriculum_id, plo_number, plo_code, category, description)
            VALUES (?,?,?,?,?)
            ON CONFLICT(curriculum_id, plo_number) DO UPDATE SET
                plo_code=excluded.plo_code,
                category=excluded.category,
                description=excluded.description
            """,
            (curriculum_id, plo_number, plo_code or str(plo_number), category, description),
        )
        row = conn.execute(
            "SELECT id FROM plos WHERE curriculum_id=? AND plo_number=?",
            (curriculum_id, plo_number),
        ).fetchone()
        return row["id"]


def get_plos(curriculum_id: int) -> list:
    with get_conn() as conn:
        rows = conn.execute(
            "SELECT * FROM plos WHERE curriculum_id=? ORDER BY plo_number",
            (curriculum_id,),
        ).fetchall()
        return [dict(r) for r in rows]


def delete_plo(plo_id: int):
    with get_conn() as conn:
        conn.execute("DELETE FROM plos WHERE id=?", (plo_id,))


def replace_plos(curriculum_id: int, plos_list: list):
    with get_conn() as conn:
        conn.execute("DELETE FROM plos WHERE curriculum_id=?", (curriculum_id,))
        for plo in plos_list:
            conn.execute(
                """INSERT INTO plos
                   (curriculum_id, plo_number, plo_code, category, description)
                   VALUES (?,?,?,?,?)""",
                (
                    curriculum_id,
                    plo["plo_number"],
                    plo.get("plo_code", "") or str(plo["plo_number"]),
                    plo.get("category", ""),
                    plo.get("description", ""),
                ),
            )


# ── YLO CRUD ──────────────────────────────────────────────────────────────────

def upsert_ylo(curriculum_id: int, year_number: int,
               title: str = "", indicators: str = "",
               assessment_methods: str = "",
               plo_mapping: list = None) -> int:
    plo_json = json.dumps(plo_mapping or [], ensure_ascii=False)
    with get_conn() as conn:
        conn.execute(
            """
            INSERT INTO ylos (curriculum_id, year_number, title,
                              indicators, assessment_methods, plo_mapping)
            VALUES (?,?,?,?,?,?)
            ON CONFLICT(curriculum_id, year_number) DO UPDATE SET
                title=excluded.title,
                indicators=excluded.indicators,
                assessment_methods=excluded.assessment_methods,
                plo_mapping=excluded.plo_mapping
            """,
            (curriculum_id, year_number, title, indicators, assessment_methods, plo_json),
        )
        row = conn.execute(
            "SELECT id FROM ylos WHERE curriculum_id=? AND year_number=?",
            (curriculum_id, year_number),
        ).fetchone()
        return row["id"]


def get_ylos(curriculum_id: int) -> list:
    with get_conn() as conn:
        rows = conn.execute(
            "SELECT * FROM ylos WHERE curriculum_id=? ORDER BY year_number",
            (curriculum_id,),
        ).fetchall()
        result = []
        for r in rows:
            d = dict(r)
            d["plo_mapping"] = json.loads(d.get("plo_mapping", "[]"))
            result.append(d)
        return result


def replace_ylos(curriculum_id: int, ylos_list: list):
    with get_conn() as conn:
        conn.execute("DELETE FROM ylos WHERE curriculum_id=?", (curriculum_id,))
        for ylo in ylos_list:
            conn.execute(
                """INSERT INTO ylos
                   (curriculum_id, year_number, title, indicators,
                    assessment_methods, plo_mapping)
                   VALUES (?,?,?,?,?,?)""",
                (
                    curriculum_id,
                    ylo["year_number"],
                    ylo.get("title", ""),
                    ylo.get("indicators", ""),
                    ylo.get("assessment_methods", ""),
                    json.dumps(ylo.get("plo_mapping", []), ensure_ascii=False),
                ),
            )


def upsert_course_clo(course_id: int, clo_number: int,
                      description: str = "", domain: str = "",
                      teaching_strategy: str = "", assessment_method: str = "",
                      pass_threshold_pct: float = 50.0,
                      plo_mapping: list = None) -> int:
    plo_json = json.dumps(plo_mapping or [], ensure_ascii=False)
    with get_conn() as conn:
        conn.execute(
            """
            INSERT INTO course_clos (
                course_id, clo_number, description, domain,
                teaching_strategy, assessment_method, pass_threshold_pct, plo_mapping
            )
            VALUES (?,?,?,?,?,?,?,?)
            ON CONFLICT(course_id, clo_number) DO UPDATE SET
                description=excluded.description,
                domain=excluded.domain,
                teaching_strategy=excluded.teaching_strategy,
                assessment_method=excluded.assessment_method,
                pass_threshold_pct=excluded.pass_threshold_pct,
                plo_mapping=excluded.plo_mapping
            """,
            (
                course_id, clo_number, description, domain,
                teaching_strategy, assessment_method, pass_threshold_pct, plo_json,
            ),
        )
        row = conn.execute(
            "SELECT id FROM course_clos WHERE course_id=? AND clo_number=?",
            (course_id, clo_number),
        ).fetchone()
        return row["id"]


def get_course_clos(course_id: int) -> list:
    with get_conn() as conn:
        rows = conn.execute(
            "SELECT * FROM course_clos WHERE course_id=? ORDER BY clo_number",
            (course_id,),
        ).fetchall()
        result = []
        for row in rows:
            data = dict(row)
            data["plo_mapping"] = json.loads(data.get("plo_mapping", "[]"))
            result.append(data)
        return result


def replace_course_clos(course_id: int, clos_list: list):
    with get_conn() as conn:
        conn.execute("DELETE FROM course_clos WHERE course_id=?", (course_id,))
        for clo in clos_list:
            conn.execute(
                """
                INSERT INTO course_clos (
                    course_id, clo_number, description, domain,
                    teaching_strategy, assessment_method, pass_threshold_pct, plo_mapping
                )
                VALUES (?,?,?,?,?,?,?,?)
                """,
                (
                    course_id,
                    clo["clo_number"],
                    clo.get("description", ""),
                    clo.get("domain", ""),
                    clo.get("teaching_strategy", ""),
                    clo.get("assessment_method", ""),
                    clo.get("pass_threshold_pct", 50.0),
                    json.dumps(clo.get("plo_mapping", []), ensure_ascii=False),
                ),
            )


def get_course_teaching_plan(course_id: int) -> list:
    with get_conn() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM course_teaching_plan WHERE course_id=? ORDER BY seq, week, id",
            (course_id,),
        ).fetchall()]


def replace_course_teaching_plan(course_id: int, plan_list: list):
    with get_conn() as conn:
        conn.execute("DELETE FROM course_teaching_plan WHERE course_id=?", (course_id,))
        for idx, item in enumerate(plan_list):
            conn.execute(
                """
                INSERT INTO course_teaching_plan (
                    course_id, seq, week, week_label, llo_text, topic,
                    activities, teaching_method, media, assessment_tools,
                    hours_planned, hours_theory, hours_practice, hours_self
                )
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    course_id,
                    item.get("seq", idx),
                    item.get("week", 0),
                    item.get("week_label", ""),
                    item.get("llo_text", ""),
                    item.get("topic", ""),
                    item.get("activities", ""),
                    item.get("teaching_method", ""),
                    item.get("media", ""),
                    item.get("assessment_tools", ""),
                    item.get("hours_planned", 0),
                    item.get("hours_theory", 0),
                    item.get("hours_practice", 0),
                    item.get("hours_self", 0),
                ),
            )


def get_course_assessments(course_id: int) -> list:
    with get_conn() as conn:
        rows = conn.execute(
            "SELECT * FROM course_assessments WHERE course_id=? ORDER BY seq, id",
            (course_id,),
        ).fetchall()
        result = []
        for row in rows:
            data = dict(row)
            data["clo_mapping"] = json.loads(data.get("clo_mapping", "[]"))
            result.append(data)
        return result


def replace_course_assessments(course_id: int, assessments_list: list):
    with get_conn() as conn:
        conn.execute("DELETE FROM course_assessments WHERE course_id=?", (course_id,))
        for idx, item in enumerate(assessments_list):
            conn.execute(
                """
                INSERT INTO course_assessments (
                    course_id, seq, name, full_score, weight_pct,
                    clo_mapping, assessment_period, eval_criteria, pass_threshold
                )
                VALUES (?,?,?,?,?,?,?,?,?)
                """,
                (
                    course_id,
                    item.get("seq", idx),
                    item.get("name", ""),
                    item.get("full_score", 100),
                    item.get("weight_pct", 0),
                    json.dumps(item.get("clo_mapping", []), ensure_ascii=False),
                    item.get("assessment_period", ""),
                    item.get("eval_criteria", ""),
                    item.get("pass_threshold", 50.0),
                ),
            )


def get_course_resources(course_id: int) -> list:
    with get_conn() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM course_resources WHERE course_id=? ORDER BY seq, id",
            (course_id,),
        ).fetchall()]


def replace_course_resources(course_id: int, resources_list: list):
    with get_conn() as conn:
        conn.execute("DELETE FROM course_resources WHERE course_id=?", (course_id,))
        for idx, item in enumerate(resources_list):
            conn.execute(
                """
                INSERT INTO course_resources
                    (course_id, seq, resource_type, citation_text, url, note)
                VALUES (?,?,?,?,?,?)
                """,
                (
                    course_id,
                    item.get("seq", idx),
                    item.get("resource_type", ""),
                    item.get("citation_text", ""),
                    item.get("url", ""),
                    item.get("note", ""),
                ),
            )


def get_tqf3_staff(tqf3_id: int) -> list:
    with get_conn() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM tqf3_staff WHERE tqf3_id=? ORDER BY role, seq, id",
            (tqf3_id,),
        ).fetchall()]


def replace_tqf3_staff(tqf3_id: int, staff_list: list):
    with get_conn() as conn:
        conn.execute("DELETE FROM tqf3_staff WHERE tqf3_id=?", (tqf3_id,))
        for idx, item in enumerate(staff_list):
            conn.execute(
                """
                INSERT INTO tqf3_staff (tqf3_id, role, seq, name)
                VALUES (?,?,?,?)
                """,
                (
                    tqf3_id,
                    item.get("role", ""),
                    item.get("seq", idx),
                    item.get("name", ""),
                ),
            )


def copy_course_template_to_tqf3(course_id: int, tqf3_id: int):
    clos = get_course_clos(course_id)
    plan_rows = get_course_teaching_plan(course_id)
    assessments = get_course_assessments(course_id)

    with get_conn() as conn:
        conn.execute("DELETE FROM clos WHERE tqf3_id=?", (tqf3_id,))
        for clo in clos:
            conn.execute(
                """
                INSERT INTO clos (
                    tqf3_id, clo_number, description,
                    teaching_strategy, assessment_method, target_pct
                )
                VALUES (?,?,?,?,?,?)
                """,
                (
                    tqf3_id,
                    clo["clo_number"],
                    clo["description"],
                    clo.get("teaching_strategy", ""),
                    clo.get("assessment_method", ""),
                    clo.get("pass_threshold_pct", 50.0),
                ),
            )

        conn.execute("DELETE FROM teaching_plan WHERE tqf3_id=?", (tqf3_id,))
        for row in plan_rows:
            conn.execute(
                """
                INSERT INTO teaching_plan (
                    tqf3_id, week, week_label, llo_text, topic,
                    activities, media, assessment_tools,
                    hours_planned, hours_theory, hours_practice, hours_self,
                    teaching_method
                )
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    tqf3_id,
                    row.get("week", 0),
                    row.get("week_label", ""),
                    row.get("llo_text", ""),
                    row.get("topic", ""),
                    row.get("activities", ""),
                    row.get("media", ""),
                    row.get("assessment_tools", ""),
                    row.get("hours_planned", 0),
                    row.get("hours_theory", 0),
                    row.get("hours_practice", 0),
                    row.get("hours_self", 0),
                    row.get("teaching_method", ""),
                ),
            )

        conn.execute("DELETE FROM assessments WHERE tqf3_id=?", (tqf3_id,))
        for item in assessments:
            conn.execute(
                """
                INSERT INTO assessments (
                    tqf3_id, name, full_score, weight_pct, clo_mapping,
                    assessment_period, eval_criteria, pass_threshold
                )
                VALUES (?,?,?,?,?,?,?,?)
                """,
                (
                    tqf3_id,
                    item.get("name", ""),
                    item.get("full_score", 100),
                    item.get("weight_pct", 0),
                    json.dumps(item.get("clo_mapping", []), ensure_ascii=False),
                    item.get("assessment_period", ""),
                    item.get("eval_criteria", ""),
                    item.get("pass_threshold", 50.0),
                ),
            )

        conn.execute("UPDATE tqf3 SET source_type='generated' WHERE id=?", (tqf3_id,))
        conn.execute(
            """
            UPDATE course_offerings
            SET source_type='generated',
                updated_at=datetime('now','localtime')
            WHERE id=(SELECT offering_id FROM tqf3 WHERE id=?)
            """,
            (tqf3_id,),
        )

    print(
        f"[DB] Copied course template -> tqf3_id={tqf3_id}: "
        f"{len(clos)} CLOs, {len(plan_rows)} plan rows, {len(assessments)} assessments"
    )


if __name__ == "__main__":
    init_db()
    print("Database initialized successfully.")
