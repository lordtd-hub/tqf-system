import os
import shutil
import tempfile
import unittest

from docx import Document

import database as db
import generate_tqf3


class GenerateTQF3Tests(unittest.TestCase):
    def setUp(self):
        self._tmpdir = tempfile.mkdtemp()
        self._orig_db_path = db.DB_PATH
        db.DB_PATH = os.path.join(self._tmpdir, "test_tqf3.db")
        db.init_db()

        self.curriculum_id = db.upsert_curriculum("69", "หลักสูตรทดสอบ", 2569)
        self.course_id = db.upsert_course(
            "MAT999",
            "วิชาทดสอบ",
            name_en="Test Course",
            credits_text="3(3-0-6)",
            credit_lecture=3,
            credit_lab=0,
            credit_self=6,
            course_type="วิชาบังคับ",
            description_th="คำอธิบายภาษาไทย",
            description_en="English description",
            curriculum_id=self.curriculum_id,
        )
        self.tqf3_id = db.upsert_tqf3(
            self.course_id,
            1,
            2569,
            instructor_main="อ.หัวหน้าวิชา",
            instructors=["อ.ผู้สอนร่วม"],
        )

    def tearDown(self):
        db.DB_PATH = self._orig_db_path
        shutil.rmtree(self._tmpdir, ignore_errors=True)

    def test_ensure_list_accepts_json_and_native_list(self):
        self.assertEqual(generate_tqf3._ensure_list("[1, 2, 3]"), [1, 2, 3])
        self.assertEqual(generate_tqf3._ensure_list([4, 5]), [4, 5])
        self.assertEqual(generate_tqf3._ensure_list(None), [])

    def test_load_data_for_tqf3_prefers_snapshot_and_falls_back_to_course_template(self):
        db.replace_plos(
            self.curriculum_id,
            [{"plo_number": 1, "plo_code": "1.1", "category": "ความรู้", "description": "PLO ทดสอบ"}],
        )
        db.replace_course_clos(
            self.course_id,
            [{
                "clo_number": 1,
                "description": "CLO ระดับวิชา",
                "domain": "ความรู้",
                "teaching_strategy": "บรรยาย",
                "assessment_method": "สอบ",
                "pass_threshold_pct": 50.0,
                "plo_mapping": [1],
            }],
        )
        db.replace_clos(
            self.tqf3_id,
            [{
                "clo_number": 1,
                "description": "CLO จาก snapshot",
                "teaching_strategy": "อภิปราย",
                "assessment_method": "งาน",
                "target_pct": 60.0,
            }],
        )
        db.replace_course_assessments(
            self.course_id,
            [{
                "name": "สอบปลายภาค",
                "full_score": 100,
                "weight_pct": 40,
                "clo_mapping": [1],
                "assessment_period": "สัปดาห์ 15",
                "eval_criteria": "ผ่าน 50%",
                "pass_threshold": 50,
            }],
        )
        db.replace_course_teaching_plan(
            self.course_id,
            [{
                "week": 1,
                "week_label": "1(3)",
                "llo_text": "LLO1",
                "topic": "บทนำ",
                "teaching_method": "บรรยาย",
                "media": "สไลด์",
                "hours_theory": 3,
            }],
        )
        db.replace_course_resources(
            self.course_id,
            [{"resource_type": "book", "citation_text": "หนังสือทดสอบ", "url": "", "note": ""}],
        )

        data = generate_tqf3.load_data_for_tqf3(tqf3_id=self.tqf3_id)

        self.assertEqual(data["course"]["code"], "MAT999")
        self.assertEqual(data["semester"], 1)
        self.assertEqual(data["year"], 2569)
        self.assertEqual(data["clos"][0]["description"], "CLO จาก snapshot")
        self.assertEqual(data["assessments"][0]["name"], "สอบปลายภาค")
        self.assertEqual(data["teaching_plan"][0]["topic"], "บทนำ")
        self.assertEqual(data["resources"][0]["citation_text"], "หนังสือทดสอบ")
        self.assertEqual(data["plos"][0]["plo_code"], "1.1")
        self.assertEqual([row["name"] for row in data["staff"]], ["อ.หัวหน้าวิชา", "อ.ผู้สอนร่วม"])

    def test_generate_tqf3_docx_can_build_from_tqf3_id_only(self):
        db.replace_plos(
            self.curriculum_id,
            [{"plo_number": 1, "plo_code": "1.1", "category": "ความรู้", "description": "PLO ทดสอบ"}],
        )
        db.replace_course_clos(
            self.course_id,
            [{
                "clo_number": 1,
                "description": "อธิบายแนวคิดพื้นฐานได้",
                "domain": "ความรู้",
                "teaching_strategy": "บรรยาย",
                "assessment_method": "สอบ",
                "pass_threshold_pct": 50.0,
                "plo_mapping": [1],
            }],
        )
        db.replace_course_assessments(
            self.course_id,
            [{
                "name": "สอบย่อย",
                "full_score": 20,
                "weight_pct": 20,
                "clo_mapping": [1],
                "assessment_period": "สัปดาห์ 4",
                "eval_criteria": "ผ่านครึ่งหนึ่ง",
                "pass_threshold": 50,
            }],
        )
        db.replace_course_teaching_plan(
            self.course_id,
            [{
                "week": 1,
                "week_label": "1(3)",
                "llo_text": "LLO1",
                "topic": "บทนำ",
                "teaching_method": "บรรยาย",
                "media": "สไลด์",
                "hours_theory": 3,
            }],
        )

        out = os.path.join(self._tmpdir, "generated_tqf3.docx")
        result = generate_tqf3.generate_tqf3_docx(tqf3_id=self.tqf3_id, output_path=out)

        self.assertEqual(result, out)
        self.assertTrue(os.path.exists(out))

        doc = Document(out)
        self.assertEqual(len(doc.tables), 6)
        self.assertEqual(doc.paragraphs[3].text.strip(), "MAT999 วิชาทดสอบ")
        self.assertIn("ภาคเรียนที่ 1 ปีการศึกษา 2569", doc.paragraphs[4].text)
        self.assertIn("MAT999", doc.tables[0].cell(2, 0).text)
        self.assertIn("วิชาทดสอบ", doc.tables[0].cell(2, 0).text)
        self.assertIn("PLO1.1", doc.tables[1].rows[1].cells[0].text)
        self.assertIn("CLO1", doc.tables[1].rows[1].cells[1].text)
        self.assertIn("บทนำ", doc.tables[2].rows[1].cells[2].text)
        self.assertIn("สัปดาห์ 4", doc.tables[4].rows[1].cells[2].text)
        self.assertIn("สอบย่อย", doc.tables[4].rows[1].cells[1].text)


if __name__ == "__main__":
    unittest.main()
