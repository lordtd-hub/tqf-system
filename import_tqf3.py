"""
import_tqf3.py — Parse มคอ.3 .docx/.pdf and import into database
"""

import re
import os
import sys
import subprocess
import tempfile
import shutil
from docx import Document
import database as db


# ════════════════════════════════════════════════
# PDF SUPPORT — mock classes ให้ทำงานเหมือน python-docx API
# ════════════════════════════════════════════════

class _Para:
    """Mock paragraph"""
    def __init__(self, text: str):
        self.text = text

class _Cell:
    """Mock cell"""
    def __init__(self, text: str):
        self.text = text
        lines = [l for l in text.split("\n") if l.strip()]
        self.paragraphs = [_Para(l) for l in lines] or [_Para("")]

class _Row:
    """Mock row"""
    def __init__(self, cells: list):
        self.cells = [_Cell(c) for c in cells]

class _Table:
    """Mock table"""
    def __init__(self, rows: list):
        self.rows = [_Row(r) for r in rows]

class PDFDoc:
    """
    Mock Document ที่สร้างจาก pdfplumber
    ทำงานเหมือน python-docx Document:
      .paragraphs → list of _Para
      .tables     → list of _Table
    """
    def __init__(self, filepath: str):
        self.paragraphs = []
        self.tables = []
        self._load(filepath)

    def _load(self, filepath: str):
        try:
            import pdfplumber
        except ImportError:
            raise ImportError(
                "ต้องติดตั้ง pdfplumber ก่อน:\n"
                "  pip install pdfplumber"
            )

        text_lines = []
        raw_tables = []

        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                # ── ดึง text ──────────────────────────
                txt = page.extract_text(x_tolerance=3, y_tolerance=3) or ""
                for line in txt.split("\n"):
                    line = line.strip()
                    if line:
                        text_lines.append(line)

                # ── ดึง tables ────────────────────────
                # ปรับ settings ให้จับ table ที่มี border บาง
                tbl_settings = {
                    "vertical_strategy": "lines",
                    "horizontal_strategy": "lines",
                    "intersection_tolerance": 5,
                }
                tables = page.extract_tables(tbl_settings)
                # fallback: text-based strategy ถ้าไม่เจอ
                if not tables:
                    tables = page.extract_tables({
                        "vertical_strategy": "text",
                        "horizontal_strategy": "text",
                        "min_words_vertical": 3,
                    })

                for tbl in (tables or []):
                    cleaned = []
                    for row in tbl:
                        cr = [" ".join(str(c).split()) if c else "" for c in row]
                        if any(cr):
                            cleaned.append(cr)
                    if len(cleaned) >= 2:
                        raw_tables.append(cleaned)

        self.paragraphs = [_Para(line) for line in text_lines]
        self.tables = [_Table(rows) for rows in raw_tables]


def open_document(filepath: str):
    """
    เปิดไฟล์ .docx/.doc/.rtf/.pdf แล้วคืน Document-like object
    """
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return PDFDoc(filepath)
    docx_path = ensure_docx(filepath)
    return Document(docx_path)


# ════════════════════════════════════════════════
# FILE CONVERSION: .doc / .rtf → .docx
# ════════════════════════════════════════════════

def ensure_docx(filepath: str) -> str:
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".docx":
        return filepath
    lo = shutil.which("libreoffice") or shutil.which("soffice")
    if not lo:
        raise RuntimeError(
            "ไม่พบ LibreOffice — กรุณาติดตั้ง LibreOffice เพื่อรองรับไฟล์ .doc/.rtf\n"
            "ดาวน์โหลดได้ที่ https://www.libreoffice.org/"
        )
    out_dir = tempfile.mkdtemp()
    result = subprocess.run(
        [lo, "--headless", "--convert-to", "docx", filepath, "--outdir", out_dir],
        capture_output=True, text=True, timeout=60
    )
    if result.returncode != 0:
        raise RuntimeError(f"แปลงไฟล์ไม่สำเร็จ: {result.stderr}")
    basename = os.path.splitext(os.path.basename(filepath))[0] + ".docx"
    converted = os.path.join(out_dir, basename)
    if not os.path.exists(converted):
        raise RuntimeError(f"ไม่พบไฟล์ที่แปลงแล้ว: {converted}")
    return converted


# ════════════════════════════════════════════════
# TEXT UTILS
# ════════════════════════════════════════════════

def clean(text: str) -> str:
    return " ".join(text.split()).strip()


def cell_text(cell) -> str:
    return clean("\n".join(p.text for p in cell.paragraphs))


def table_to_rows(table):
    return [[cell_text(c) for c in row.cells] for row in table.rows]


# ════════════════════════════════════════════════
# EXTRACT SECTIONS
# ════════════════════════════════════════════════

def get_merged_cell_text(table):
    """
    ตาราง มคอ.3 มักเป็น single-column merged cells
    คืน list ของ string (เนื้อหาแต่ละแถว, de-duplicated)
    """
    lines = []
    for row in table.rows:
        texts = [clean(c.text) for c in row.cells]
        # de-duplicate merged cells
        unique = list(dict.fromkeys(t for t in texts if t))
        merged = " ".join(unique)
        if merged:
            lines.append(merged)
    return lines


def extract_header_info(doc):
    """
    ดึงข้อมูล หมวด 1 จากตารางแรก (ตารางกล่อง info — merged cells)
    คืน dict: code, name_th, name_en, credits_text, semester, year,
              faculty, department, prerequisite, course_type,
              instructor_main, instructors, location
    """
    info = {
        "code": "", "name_th": "", "name_en": "",
        "credits_text": "", "credit_lecture": 0, "credit_lab": 0, "credit_self": 0,
        "semester": 1, "year": 2568,
        "faculty": "", "department": "",
        "prerequisite": "ไม่มี", "course_type": "",
        "description_th": "", "description_en": "",
        "instructor_main": "", "instructors": [], "location": "",
        "objectives": ""
    }

    # ── ดึงจาก title paragraphs (ก่อนตาราง) ────
    full_text = "\n".join(p.text for p in doc.paragraphs[:30])

    # รหัส + ชื่อวิชาจาก title เช่น "SMA0901 สัมมนาคณิตศาสตร์"
    for para in doc.paragraphs[:10]:
        txt = para.text.strip()
        m_title = re.match(r"([A-Z]{2,6}\d{3,6}[A-Z]?)\s+(.+)", txt)
        if m_title:
            info["code"] = m_title.group(1)
            info["name_th"] = m_title.group(2)
            break

    # ภาคเรียนที่ X ปีการศึกษา YYYY
    m = re.search(r"ภาคเรียนที่\s*(\d)\s*ปีการศึกษา\s*(\d{4})", full_text)
    if m:
        info["semester"] = int(m.group(1))
        info["year"] = int(m.group(2))

    # ── parse ตารางแรก (info box) ────────────────
    for table in doc.tables[:3]:
        lines = get_merged_cell_text(table)
        all_text = " ".join(lines)

        # ตรวจว่าเป็นตาราง info box
        if not any(kw in all_text for kw in ["รหัสวิชา", "หน่วยกิต", "คณะ"]):
            continue

        for line in lines:
            # คณะ/สาขา
            if "คณะ/สาขา" in line:
                txt = re.sub(r".*คณะ/สาขา/วิชาเอก\s*", "", line).strip()
                # รูปแบบ "คณะ... สาขาวิชา..."
                m_fac = re.search(r"(คณะ[^\s]+(?:\s+\S+)?)\s+(สาขาวิชา.+)", txt)
                if m_fac:
                    info["faculty"]    = m_fac.group(1)
                    info["department"] = m_fac.group(2)
                else:
                    info["faculty"] = txt

            # รหัสวิชา + ชื่อ
            if "รหัสวิชา" in line:
                m2 = re.search(r"รหัสวิชา\s*([A-Z]{2,6}\d{3,6}[A-Z]?)", line)
                if m2:
                    info["code"] = m2.group(1)
                m3 = re.search(r"ชื่อวิชา\s*\(ไทย\)\s*(.+?)(?:ชื่อวิชา\s*\(อังกฤษ\)|$)", line)
                if m3:
                    info["name_th"] = clean(m3.group(1))
                m4 = re.search(r"ชื่อวิชา\s*\(อังกฤษ\)\s*(.+?)$", line)
                if m4:
                    info["name_en"] = clean(m4.group(1))

            # หน่วยกิต
            if "หน่วยกิต" in line:
                m5 = re.search(r"(\d+)\s*\(\s*(\d+)\s*-\s*(\d+)\s*-\s*(\d+)\s*\)", line)
                if m5:
                    info["credits_text"] = m5.group(0).replace(" ", "")
                    info["credit_lecture"] = int(m5.group(2))
                    info["credit_lab"]     = int(m5.group(3))
                    info["credit_self"]    = int(m5.group(4))

            # วิชาที่ต้องเรียนก่อน / ประเภทวิชา
            if "รายวิชาที่ต้องเรียนก่อน" in line:
                info["prerequisite"] = "ไม่มี" if "ไม่มี" in line else (
                    re.sub(r".*รายวิชาที่ต้องเรียนก่อน.*?(\(ถ้ามี\))?\s*", "", line).strip()
                    or "ไม่มี")
            if "ประเภทของรายวิชา" in line:
                # บางฟอร์มมี "ชื่อหลักสูตร...ประเภทของรายวิชา ชื่อ... ประเภทของรายวิชา ค่าจริง"
                # → ต้องข้าม match แรกที่ค่าขึ้นต้นด้วย "ชื่อหลักสูตร"
                # ใช้ negative lookahead เพื่อจับเฉพาะ match ที่ค่าไม่ขึ้นต้นด้วย "ชื่อ"
                for m6 in re.finditer(r"ประเภทของรายวิชา\s+((?!ชื่อ)\S[^\t\n]*?)(?:\s+(?:รายวิชาที่|ชื่อหลักสูตร|ภาษา|อาจารย์|\d+\.)|$)", line):
                    ct = m6.group(1).strip()
                    ct = re.sub(r"\s*(รายวิชาที่|ภาษา|อาจารย์|ชื่อหลักสูตร|\d+\.).*$", "", ct).strip()
                    if ct and len(ct) > 1:
                        info["course_type"] = ct
                        break

            # คำอธิบายรายวิชา
            if "คำอธิบายรายวิชา" in line:
                # เก็บข้อความหลัง keyword
                desc = re.sub(r".*คำอธิบายรายวิชา\s*(ภาษาไทย\s*)?", "", line).strip()
                if desc and len(desc) > 10:
                    info["description_th"] = desc[:300]

            # อาจารย์ผู้รับผิดชอบ
            if "อาจารย์ผู้รับผิดชอบ" in line and "อาจารย์ผู้สอน" not in line:
                txt = re.sub(r".*อาจารย์ผู้รับผิดชอบ\s*", "", line).strip()
                if txt:
                    info["instructor_main"] = txt

            # สถานที่เรียน
            if re.search(r"สถานที่เรียน|สถานที่จัดการเรียน", line):
                txt = re.sub(r".*สถานที่เรียน\s*|.*สถานที่จัดการเรียนการสอน\s*", "", line).strip()
                if txt:
                    info["location"] = txt

            # ภาคการศึกษา/ปีการศึกษาในตาราง
            if "ภาคการศึกษา" in line and "ปีการศึกษา" in line:
                m7 = re.search(r"(\d)/(\d{4})", line)
                if m7:
                    info["semester"] = int(m7.group(1))
                    info["year"] = int(m7.group(2))

        break  # ใช้ตารางแรกที่ตรงเงื่อนไข

    # ── parse หมวด 9 จาก paragraphs (อาจารย์ผู้สอน/ผู้รับผิดชอบ) ──────────
    # seminar format เก็บอาจารย์ไว้ใน paragraph ไม่ใช่ใน table
    in_91 = False  # คณะกรรมการบริหารรายวิชา
    in_92 = False  # อาจารย์ผู้สอนรายวิชา
    in_93 = False  # สถานที่จัดการเรียนการสอน

    for para in doc.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue

        # หัวข้อ 9.1 / 9.2 / 9.3
        if re.search(r'9[._\s]*1', txt) and ('คณะกรรมการ' in txt or 'บริหาร' in txt):
            in_91, in_92, in_93 = True, False, False
            continue
        if re.search(r'9[._\s]*2', txt) and 'อาจารย์ผู้สอน' in txt:
            in_91, in_92, in_93 = False, True, False
            continue
        if re.search(r'9[._\s]*3', txt) and 'สถานที่' in txt:
            in_91, in_92, in_93 = False, False, True
            continue
        # หัวข้อถัดไป (หมวด 10+) — หยุด
        if re.match(r'^(?:หมวดที่\s*)?\d{2,}|^10[._\s]', txt):
            in_91 = in_92 = in_93 = False

        if in_91:
            name = re.sub(r'^\d+[.)]\s*', '', txt).strip()
            if name and len(name) > 2:
                # เก็บคนแรกใน 9.1 เป็น instructor_main
                if not info["instructor_main"]:
                    info["instructor_main"] = name
                in_91 = False  # แค่คนแรก

        elif in_92:
            name = re.sub(r'^\d+[.)]\s*', '', txt).strip()
            if name and len(name) > 2 and name not in info["instructors"]:
                info["instructors"].append(name)
                # ถ้ายังไม่มี instructor_main ให้ใช้คนแรกใน 9.2
                if not info["instructor_main"]:
                    info["instructor_main"] = name

        elif in_93:
            loc = re.sub(r'^\d+[.)]\s*', '', txt).strip()
            if loc and len(loc) > 2 and not info["location"]:
                info["location"] = loc

    return info


def extract_clos(doc):
    """
    ดึง CLOs
    รองรับ 2 รูปแบบ:
    1. ตาราง CLO matrix (Table 1 ใน seminar format — first column = CLO description)
    2. ตาราง CLO+กลยุทธ์ (classic format)
    """
    clos = []
    clo_counter = 0

    for table in doc.tables:
        rows = table_to_rows(table)
        if not rows:
            continue
        header_str = " ".join(rows[0])

        # ── รูปแบบ 1: ตาราง Matrix PLO/CLO ──
        # header มี "จุดมุ่งหมายของรายวิชา" หรือ "ผลลัพธ์การเรียนรู้รายวิชา"
        # CLO อยู่ใน column แรก แต่ละ row ที่ไม่ใช่ header
        if any(kw in header_str for kw in ["จุดมุ่งหมายของรายวิชา", "CLO", "ผลลัพธ์การเรียนรู้รายวิชา"]):
            # header rows คือแถวที่ cell แรกซ้ำกับ header หรือมี PLO/คุณธรรม
            HEADER_KEYWORDS = {"จุดมุ่งหมายของรายวิชา", "ผลลัพธ์การเรียนรู้",
                               "คุณธรรม", "ความรู้", "ทักษะ", "PLO", "CLO"}
            for row in rows:
                first_cell = row[0] if row else ""
                if not first_cell or len(first_cell) < 8:
                    continue
                # ข้าม header rows
                if any(kw in first_cell for kw in HEADER_KEYWORDS):
                    continue
                # ข้ามแถวที่เป็นตัวเลข/รหัส PLO เช่น "1.1", "2.3"
                if re.match(r"^\d+\.\d+$", first_cell.strip()):
                    continue

                # รูปแบบ "N. ข้อความ" หรือ "CLO N ข้อความ"
                m_clo = re.match(r"^(?:CLO\s*)?(\d+)[.)]\s*(.{8,})", first_cell, re.UNICODE)
                if m_clo:
                    clo_counter += 1
                    clos.append({
                        "clo_number": int(m_clo.group(1)),
                        "description": m_clo.group(2).strip(),
                        "teaching_strategy": "",
                        "assessment_method": "",
                        "indicator": "",
                        "target_pct": 50.0,
                    })
                elif len(first_cell) >= 15:
                    # ไม่มีเลขนำหน้า — นับเองตามลำดับ (format ของ complex/seminar)
                    clo_counter += 1
                    clos.append({
                        "clo_number": clo_counter,
                        "description": first_cell.strip(),
                        "teaching_strategy": "",
                        "assessment_method": "",
                        "indicator": "",
                        "target_pct": 50.0,
                    })
            if clos:
                break

        # ── รูปแบบ 2a: PLO | CLO | กลยุทธ์ | ประเมิน (Intro to AI format) ──
        # col 0 = PLO ref, col 1 = CLO description
        PLO_REF = re.compile(r"^PLO\d+", re.IGNORECASE)
        has_plo_col = any(kw in header_str for kw in
                          ["ผลการเรียนรู้ที่คาดหวังของหลักสูตร", "PLO"])
        has_clo_col = any(kw in header_str for kw in
                          ["ผลการเรียนรู้ที่คาดหวังของรายวิชา", "CLO"])
        if has_plo_col and has_clo_col and not clos:
            for row in rows[1:]:
                if len(row) < 2:
                    continue
                clo_cell = row[1]   # CLO อยู่ใน col 1
                if not clo_cell or len(clo_cell) < 5:
                    continue
                # ข้าม header ซ้ำ
                if any(kw in clo_cell for kw in ["ผลการเรียนรู้", "CLO"]) and len(clo_cell) < 20:
                    continue
                # strip "CLO1 " prefix → เอาเฉพาะ description
                m_clo = re.match(r"^CLO\s*(\d+)\s+(.{5,})", clo_cell, re.DOTALL)
                if m_clo:
                    num  = int(m_clo.group(1))
                    desc = clean(m_clo.group(2))
                else:
                    clo_counter += 1
                    num  = clo_counter
                    desc = clean(clo_cell)
                clos.append({
                    "clo_number":       num,
                    "description":      desc,
                    "teaching_strategy": clean(row[2]) if len(row) > 2 else "",
                    "assessment_method": clean(row[3]) if len(row) > 3 else "",
                    "indicator":        "",
                    "target_pct":       50.0,
                })
            if clos:
                break

        # ── รูปแบบ 2b: CLO | กลยุทธ์ | ประเมิน (classic) ──
        if any(kw in header_str for kw in ["กลยุทธ์การสอน", "วิธีการประเมิน", "ผลการเรียนรู้ที่คาดหวัง"]) and not clos:
            for row in rows[1:]:
                if len(row) < 2:
                    continue
                desc = row[0]
                if not desc:
                    continue
                desc_clean = re.sub(r"^(CLO\s*\d+\s*|[\d]+[.)]\s*)", "", desc).strip()
                if len(desc_clean) < 5:
                    continue
                clo_counter += 1
                clos.append({
                    "clo_number": clo_counter,
                    "description": desc_clean,
                    "teaching_strategy": row[1] if len(row) > 1 else "",
                    "assessment_method": row[2] if len(row) > 2 else "",
                    "indicator": "",
                    "target_pct": 50.0,
                })
            if clos:
                break

    # ── fallback: จาก paragraphs ──────────────────
    if not clos:
        for para in doc.paragraphs:
            txt = para.text.strip()
            m = re.match(r"^(?:CLO\s*)?(\d+)[.)]\s*(.{10,})", txt)
            if m and int(m.group(1)) <= 20:
                clo_counter += 1
                clos.append({
                    "clo_number": int(m.group(1)),
                    "description": m.group(2).strip(),
                    "teaching_strategy": "",
                    "assessment_method": "",
                    "indicator": "",
                    "target_pct": 50.0,
                })

    return clos


def extract_teaching_plan(doc):
    """
    ดึงแผนการสอนรายสัปดาห์จากตารางหมวด 5
    คืน list ของ dict

    รูปแบบตาราง:
      Col 0: สัปดาห์ที่ (จำนวนชั่วโมง) เช่น "1(3)" หรือ "4-6(9)"
      Col 1: ผลลัพธ์การเรียนรู้ / CLO reference  เช่น "CLO1, CLO2"
      Col 2: หัวข้อบรรยาย (topic จริง) เช่น "บทนำ, ลิมิต..."
      Col 3: วิธีการสอน
      ...

    บางรูปแบบ topic อยู่ใน col 1 และไม่มี CLO column
    """
    plan = []

    CLO_PATTERN = re.compile(r"^(CLO\d+|PLO\d+|ALL\s*CLO)", re.IGNORECASE)

    for table in doc.tables:
        rows = table_to_rows(table)
        if not rows:
            continue
        header = " ".join(rows[0])
        if not any(kw in header for kw in ["สัปดาห์", "หัวข้อ", "จำนวนชั่วโมง"]):
            continue

        # ตรวจว่า col 1 เป็น CLO/ผลลัพธ์การเรียนรู้ column ไหม
        # ดูจาก header: ถ้า col 1 header มีคำว่า "ผลลัพธ์" / "CLO" → topic อยู่ที่ col 2
        header_row = rows[0] if rows else []
        col1_header = header_row[1] if len(header_row) > 1 else ""
        col1_is_clo = any(kw in col1_header for kw in ["ผลลัพธ์", "CLO", "PLO"])
        if not col1_is_clo:
            # fallback: ตรวจ content ของแถวแรก
            sample_col1 = [rows[i][1] for i in range(1, min(4, len(rows))) if len(rows[i]) > 1]
            col1_is_clo = any(CLO_PATTERN.match(c) for c in sample_col1)

        for row in rows[1:]:
            if len(row) < 2:
                continue

            # ── Week number & hours ────────────────────────────────────────
            week_txt = row[0] if row else ""
            # หา week จาก "1(3)", "4-6(9)", "1-3" ฯลฯ
            m_week = re.search(r"(\d+)", week_txt)
            week = int(m_week.group(1)) if m_week else len(plan) + 1
            # หา hours จาก "(3)" หรือ "(9)" ใน col 0
            m_hrs = re.search(r"\((\d+(?:\.\d+)?)\)", week_txt)
            hours = float(m_hrs.group(1)) if m_hrs else 0.0

            # ── Topic ─────────────────────────────────────────────────────
            if col1_is_clo and len(row) > 2:
                topic = row[2]   # actual topic is in col 2
            else:
                topic = row[1]   # topic is in col 1

            if not topic or topic == week_txt:
                continue

            # ── Fallback: hours ยังเป็น 0 → ลองหาใน column อื่น ─────────
            if hours == 0:
                for col in row[1:]:
                    m2 = re.search(r"\b(\d+(?:\.\d+)?)\b", col)
                    if m2 and float(m2.group(1)) <= 50:  # ไม่ใช่ปีการศึกษา
                        hours = float(m2.group(1))
                        break

            teaching_method = row[3] if (col1_is_clo and len(row) > 3) else (row[-1] if len(row) > 2 else "")

            plan.append({
                "week": week,
                "topic": topic,
                "hours_planned": hours,
                "teaching_method": teaching_method,
            })

        if plan:
            break

    return plan


def extract_assessments(doc):
    """
    ดึงแผนการประเมินผลจากตารางหมวด 6
    คืน list ของ dict
    """
    assessments = []

    for table in doc.tables:
        rows = table_to_rows(table)
        if not rows:
            continue
        header = " ".join(rows[0])
        # header รูปแบบจริง: "รายการ | ร้อยละ" หรือ keyword อื่น
        if not any(kw in header for kw in
                   ["กิจกรรมการประเมิน", "สัดส่วน", "น้ำหนัก", "คะแนน", "ร้อยละ", "รายการ"]):
            continue
        # ต้องมีคอลัมน์ที่เป็น % (มีตัวเลขใน col ที่ 2 ขึ้นไป)
        has_pct = any(
            re.search(r"\d+", c) for row in rows[1:3] for c in (row[1:] if len(row) > 1 else [])
        )
        if not has_pct:
            continue

        for row in rows[1:]:
            if len(row) < 2:
                continue
            name = row[0].strip()
            # ข้ามแถว header ซ้ำ, แถวว่าง, และแถว "รวม"
            if not name or name.lower() in {"รวม", "total", "รายการ"}:
                continue
            # แถวที่ชื่อเป็นเลขเปล่า เช่น "1." → ใช้ชื่อ placeholder
            if re.match(r"^\d+[.)]*$", name):
                name = f"รายการที่ {name.rstrip('.')}"
            # หา % จาก column — ให้ความสำคัญกับ cell ที่มี "%" ก่อน
            weight = 0
            # pass 1: หา cell ที่มีเครื่องหมาย % ชัดเจน
            for col in row[1:]:
                m = re.search(r"(\d+(?:\.\d+)?)\s*%", col)
                if m:
                    w = float(m.group(1))
                    if 0 < w <= 100:
                        weight = w
                        break
            # pass 2 (fallback): หาตัวเลขจาก column ขวาสุดไปซ้าย
            if weight == 0:
                for col in reversed(row[1:]):
                    m = re.search(r"(\d+(?:\.\d+)?)", col)
                    if m:
                        w = float(m.group(1))
                        if 0 < w <= 100:
                            weight = w
                            break
            if weight > 0:
                assessments.append({
                    "name": name,
                    "weight_pct": weight,
                    "clo_mapping": [],
                })

        if assessments:
            break

    return assessments


# ════════════════════════════════════════════════
# MAIN IMPORT FUNCTION
# ════════════════════════════════════════════════

def import_tqf3(filepath: str) -> dict:
    """
    Import มคอ.3 จาก .docx เข้า database
    คืน dict สรุปผล
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"ไม่พบไฟล์: {filepath}")

    print(f"[TQF3] Reading: {filepath}")
    doc = open_document(filepath)

    # ── extract ──────────────────────────────────
    info      = extract_header_info(doc)
    clos      = extract_clos(doc)
    plan      = extract_teaching_plan(doc)
    assess    = extract_assessments(doc)

    print(f"  รหัสวิชา   : {info['code']} — {info['name_th']}")
    print(f"  ภาค/ปี    : {info['semester']}/{info['year']}")
    print(f"  CLOs       : {len(clos)} ข้อ")
    print(f"  แผนการสอน  : {len(plan)} สัปดาห์")
    print(f"  การประเมิน : {len(assess)} รายการ")

    if not info["code"]:
        raise ValueError("ไม่พบรหัสวิชาในไฟล์ — กรุณาตรวจสอบรูปแบบไฟล์")

    # ── save to DB ──────────────────────────────
    db.init_db()

    course_id = db.upsert_course(
        code=info["code"],
        name_th=info["name_th"],
        name_en=info["name_en"],
        credits_text=info["credits_text"],
        credit_lecture=info["credit_lecture"],
        credit_lab=info["credit_lab"],
        credit_self=info["credit_self"],
        prerequisite=info["prerequisite"],
        course_type=info["course_type"],
        description_th=info["description_th"],
        description_en=info["description_en"],
        faculty=info["faculty"],
        department=info["department"],
    )

    tqf3_id = db.upsert_tqf3(
        course_id=course_id,
        semester=info["semester"],
        year=info["year"],
        instructor_main=info["instructor_main"],
        instructors=info["instructors"],
        location=info["location"],
        objectives=info["objectives"],
        source_file=filepath,
    )

    if clos:
        db.replace_clos(tqf3_id, clos)
    if plan:
        db.replace_teaching_plan(tqf3_id, plan)
    if assess:
        db.replace_assessments(tqf3_id, assess)

    print(f"[TQF3] Saved → course_id={course_id}, tqf3_id={tqf3_id}")

    return {
        "success": True,
        "course_id": course_id,
        "tqf3_id": tqf3_id,
        "code": info["code"],
        "name_th": info["name_th"],
        "semester": info["semester"],
        "year": info["year"],
        "clos_count": len(clos),
        "plan_weeks": len(plan),
        "assessments_count": len(assess),
    }


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python import_tqf3.py <path_to_tqf3.docx>")
        sys.exit(1)
    result = import_tqf3(sys.argv[1])
    print("\nResult:", result)
